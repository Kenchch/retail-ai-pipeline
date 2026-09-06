"""Refresh the existing editable workshop from reports/doc_figures.json.

Requires python-pptx. Preserves the six-slide template and run formatting.
Optionally build a PDF from six reviewed PNG renders with --rendered-slides DIR
(requires Pillow). Render the refreshed PPTX before using that option.
"""

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData

ROOT = Path(__file__).resolve().parents[1]


def build() -> Path:
    figures = json.loads((ROOT / "reports/doc_figures.json").read_text())
    lantern, spade = figures["pairs"]["lantern"], figures["pairs"]["spade"]
    path = ROOT / "workshop/ai_literacy_workshop.pptx"
    deck = Presentation(path)
    assert len(deck.slides) == 6, "Expected the existing six-slide workshop"
    # The source chart contains literal cached values without an embedded workbook.
    # Preserve those exact categories/values and attach the editable workbook.
    for shape in deck.slides[5].shapes:
        if shape.has_chart:
            chart = shape.chart
            data = CategoryChartData()
            data.categories = [category.label for category in chart.plots[0].categories]
            for series in chart.series:
                data.add_series(series.name, tuple(series.values))
            chart.replace_data(data)

    def text(slide, shape, value, paragraph=0, run=0):
        deck.slides[slide - 1].shapes[shape].text_frame.paragraphs[paragraph].runs[
            run
        ].text = value

    text(2, 1, "A recommendation from the live table, computed from the retail dataset")
    text(
        2,
        4,
        f"{lantern['baskets']} baskets, confidence {lantern['confidence']:.2f}, lift {lantern['lift']:.2f}",
    )
    text(
        2,
        5,
        f" the same t-light holder is recommended for {figures['t_light_products']} products",
        run=1,
    )
    text(3, 5, str(spade["baskets"]))
    text(3, 9, f"{spade['confidence']:.2f}")
    text(
        3,
        11,
        f"Of baskets with the pink spade, {spade['confidence']:.0%} also had the blue one.",
    )
    text(3, 13, f"{spade['lift']:.0f}")
    text(
        3,
        15,
        f"{spade['lift']:.0f}× the independent-purchase expectation. This is the ranking column.",
    )
    text(
        3,
        16,
        f"Back to the first row: lift {lantern['lift']:.2f} against {spade['lift']:.0f} here. Same catalogue, different strengths of association.",
    )
    text(6, 0, "Reading the simulated adoption report")
    text(6, 4, "Illustrative workshop dates")
    text(6, 5, "These changes are generated examples.")
    text(6, 5, "They do not measure workshop impact.", paragraph=1)
    text(
        6,
        5,
        "With real telemetry, check sustained use, missing data and a comparison group before attributing change to training.",
        paragraph=3,
    )
    text(
        6,
        6,
        "Exercise: identify the evidence needed to distinguish sustained use from a short-lived spike.",
    )
    deck.save(path)
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rendered-slides", type=Path)
    args = parser.parse_args()
    path = build()
    if args.rendered_slides:
        from PIL import Image

        images = [
            Image.open(args.rendered_slides / f"slide-{i}.png").convert("RGB")
            for i in range(1, 7)
        ]
        images[0].save(
            path.with_suffix(".pdf"),
            save_all=True,
            append_images=images[1:],
            resolution=144.0,
        )
        for image in images:
            image.close()
